"""Generates old.docx / new.docx: a small, non-confidential pair exercising every
kind of change the viewer shows. Run: ~/.claude/python/bin/python fixtures/make.py"""
import io
import struct
import zlib
from pathlib import Path

from docx import Document
from docx.shared import Inches


def png(width, height, rgb):
    """A solid-color PNG, built by hand so no imaging library is needed."""
    def chunk(tag, data):
        return struct.pack('>I', len(data)) + tag + data + struct.pack('>I', zlib.crc32(tag + data))
    row = b'\x00' + bytes(rgb) * width
    return (b'\x89PNG\r\n\x1a\n'
            + chunk(b'IHDR', struct.pack('>IIBBBBB', width, height, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(row * height))
            + chunk(b'IEND', b''))


def build(new: bool):
    d = Document()
    d.add_heading('Garden Plan', 0)
    d.add_paragraph('This plan describes the vegetable garden for the coming season.')
    d.add_heading('1. Beds', 1)
    d.add_paragraph(
        'The garden has four raised beds along the south fence.' if not new
        else 'The garden has five raised beds along the south fence.')
    for item in (['Tomatoes', 'Beans', 'Carrots', 'Lettuce'] if not new
                 else ['Tomatoes', 'Beans', 'Squash', 'Carrots', 'Lettuce']):
        d.add_paragraph(item, style='List Number')
    if not new:
        d.add_paragraph('The compost heap stays in the north corner.')
    d.add_heading('2. Schedule', 1)
    table = d.add_table(rows=1, cols=2)
    table.style = 'Light Grid Accent 1'
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = 'Task', 'Month'
    for task, month in [('Sow beans', 'May'), ('Plant tomatoes', 'May' if not new else 'June'), ('Harvest', 'September')]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = task, month
    d.add_heading('3. Layout', 1)
    d.add_picture(io.BytesIO(png(120, 40, (90, 160, 90))), width=Inches(3))
    if new:
        d.add_paragraph('A drip line waters all beds from the rain barrel.')
    paths = d.add_paragraph('Paths between beds are covered with ')
    paths.add_run('wood chips').bold = new  # formatting-only change
    paths.add_run('.')
    return d


here = Path(__file__).parent
build(False).save(here / 'old.docx')
build(True).save(here / 'new.docx')
print('wrote', here / 'old.docx', here / 'new.docx')
